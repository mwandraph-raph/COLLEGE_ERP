"""
XORADEX EDUCORE ERP
ENROLLMENT SUMMARY EXPORTS

Provides Excel and PDF exports for the Enrollment Summary report.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)


def create_enrollment_excel(summary):
    """
    Create an Excel workbook containing the complete
    Enrollment Summary.
    """

    workbook = Workbook()

    # ==========================================================
    # SUMMARY SHEET
    # ==========================================================

    sheet = workbook.active
    sheet.title = "Summary"

    sheet["A1"] = "XORADEX EDUCORE"
    sheet["A1"].font = Font(size=16, bold=True)

    sheet["A2"] = "ENROLLMENT SUMMARY"
    sheet["A2"].font = Font(size=13, bold=True)

    sheet["A4"] = "Metric"
    sheet["B4"] = "Value"

    for cell in sheet[4]:
        cell.font = Font(bold=True)

    summary_rows = [
        ("Total Students", summary["total_students"]),
        ("Active Students", summary["active_students"]),
        ("Inactive Students", summary["inactive_students"]),
        ("Departments", summary["department_count"]),
        ("Courses", summary["course_count"]),
        ("Programmes", summary["programme_count"]),
    ]

    row_number = 5

    for label, value in summary_rows:
        sheet.cell(row=row_number, column=1, value=label)
        sheet.cell(row=row_number, column=2, value=value)
        row_number += 1

    sheet.column_dimensions["A"].width = 25
    sheet.column_dimensions["B"].width = 18

    # ==========================================================
    # DEPARTMENT SHEET
    # ==========================================================

    sheet = workbook.create_sheet("By Department")

    headers = [
        "Department Code",
        "Department",
        "Students",
    ]

    _write_headers(sheet, headers)

    for department in summary["departments"]:
        sheet.append([
            department["programme__course__department__code"],
            department["programme__course__department__name"],
            department["total"],
        ])

    _auto_width(sheet)

    # ==========================================================
    # COURSE SHEET
    # ==========================================================

    sheet = workbook.create_sheet("By Course")

    headers = [
        "Course Code",
        "Course",
        "Students",
    ]

    _write_headers(sheet, headers)

    for course in summary["courses"]:
        sheet.append([
            course["programme__course__code"],
            course["programme__course__name"],
            course["total"],
        ])

    _auto_width(sheet)

    # ==========================================================
    # PROGRAMME SHEET
    # ==========================================================

    sheet = workbook.create_sheet("By Programme")

    headers = [
        "Programme Code",
        "Programme",
        "Award",
        "Students",
    ]

    _write_headers(sheet, headers)

    for programme in summary["programmes"]:
        sheet.append([
            programme["programme__code"],
            programme["programme__name"],
            programme["programme__award"],
            programme["total"],
        ])

    _auto_width(sheet)

    # ==========================================================
    # AWARD SHEET
    # ==========================================================

    sheet = workbook.create_sheet("By Award")

    headers = [
        "Award",
        "Students",
    ]

    _write_headers(sheet, headers)

    for award in summary["awards"]:
        sheet.append([
            award["programme__award"],
            award["total"],
        ])

    _auto_width(sheet)

    # ==========================================================
    # GENDER SHEET
    # ==========================================================

    sheet = workbook.create_sheet("By Gender")

    headers = [
        "Gender",
        "Students",
    ]

    _write_headers(sheet, headers)

    for item in summary["gender"]:
        sheet.append([
            item["gender"] or "Not specified",
            item["total"],
        ])

    _auto_width(sheet)

    # ==========================================================
    # STATUS SHEET
    # ==========================================================

    sheet = workbook.create_sheet("By Status")

    headers = [
        "Status",
        "Students",
    ]

    _write_headers(sheet, headers)

    for item in summary["status"]:
        sheet.append([
            item["status"],
            item["total"],
        ])

    _auto_width(sheet)

    # ==========================================================
    # SAVE WORKBOOK
    # ==========================================================

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    return buffer


def _write_headers(sheet, headers):
    """
    Write and format worksheet headers.
    """

    sheet.append(headers)

    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")


def _auto_width(sheet):
    """
    Automatically size worksheet columns.
    """

    for column_cells in sheet.columns:
        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)

            if len(value) > max_length:
                max_length = len(value)

        sheet.column_dimensions[column_letter].width = min(
            max_length + 3,
            60,
        )


def create_enrollment_pdf(summary):
    """
    Create a print-ready PDF Enrollment Summary.
    """

    buffer = BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "EnrollmentTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontSize=18,
        leading=22,
        spaceAfter=4,
    )

    subtitle_style = ParagraphStyle(
        "EnrollmentSubtitle",
        parent=styles["Normal"],
        alignment=TA_CENTER,
        fontSize=10,
        leading=13,
        spaceAfter=12,
    )

    heading_style = ParagraphStyle(
        "EnrollmentHeading",
        parent=styles["Heading2"],
        fontSize=12,
        leading=15,
        spaceBefore=10,
        spaceAfter=6,
    )

    normal_style = ParagraphStyle(
        "EnrollmentNormal",
        parent=styles["Normal"],
        fontSize=9,
        leading=11,
    )

    story = []

    story.append(
        Paragraph(
            "XORADEX EDUCORE",
            title_style,
        )
    )

    story.append(
        Paragraph(
            "ENROLLMENT SUMMARY",
            title_style,
        )
    )

    story.append(
        Paragraph(
            "Student enrollment overview across the academic structure.",
            subtitle_style,
        )
    )

    # ==========================================================
    # OVERVIEW
    # ==========================================================

    story.append(
        Paragraph(
            "Enrollment Overview",
            heading_style,
        )
    )

    overview_data = [
        [
            "Total Students",
            "Active",
            "Inactive",
            "Departments",
            "Courses",
            "Programmes",
        ],
        [
            str(summary["total_students"]),
            str(summary["active_students"]),
            str(summary["inactive_students"]),
            str(summary["department_count"]),
            str(summary["course_count"]),
            str(summary["programme_count"]),
        ],
    ]

    overview_table = Table(
        overview_data,
        colWidths=[
            38 * mm,
            32 * mm,
            32 * mm,
            32 * mm,
            28 * mm,
            32 * mm,
        ],
    )

    overview_table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ])
    )

    story.append(overview_table)
    story.append(Spacer(1, 8))

    # ==========================================================
    # DEPARTMENT
    # ==========================================================

    story.append(
        Paragraph(
            "Enrollment by Department",
            heading_style,
        )
    )

    department_data = [
        [
            "Department Code",
            "Department",
            "Students",
        ]
    ]

    for department in summary["departments"]:
        department_data.append([
            department["programme__course__department__code"],
            department["programme__course__department__name"],
            str(department["total"]),
        ])

    if len(department_data) == 1:
        department_data.append([
            "—",
            "No enrollment data available",
            "0",
        ])

    department_table = Table(
        department_data,
        colWidths=[40 * mm, 150 * mm, 30 * mm],
    )

    department_table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ALIGN", (2, 1), (2, -1), "RIGHT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ])
    )

    story.append(department_table)

    # ==========================================================
    # COURSE
    # ==========================================================

    story.append(
        Paragraph(
            "Enrollment by Course",
            heading_style,
        )
    )

    course_data = [
        [
            "Course Code",
            "Course",
            "Students",
        ]
    ]

    for course in summary["courses"]:
        course_data.append([
            course["programme__course__code"],
            course["programme__course__name"],
            str(course["total"]),
        ])

    if len(course_data) == 1:
        course_data.append([
            "—",
            "No course enrollment data available",
            "0",
        ])

    course_table = Table(
        course_data,
        colWidths=[40 * mm, 150 * mm, 30 * mm],
    )

    course_table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ALIGN", (2, 1), (2, -1), "RIGHT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ])
    )

    story.append(course_table)

    # ==========================================================
    # PROGRAMME
    # ==========================================================

    story.append(
        Paragraph(
            "Enrollment by Programme",
            heading_style,
        )
    )

    programme_data = [
        [
            "Programme Code",
            "Programme",
            "Award",
            "Students",
        ]
    ]

    for programme in summary["programmes"]:
        programme_data.append([
            programme["programme__code"],
            programme["programme__name"],
            str(programme["programme__award"]).replace("_", " ").title(),
            str(programme["total"]),
        ])

    if len(programme_data) == 1:
        programme_data.append([
            "—",
            "No programme enrollment data available",
            "—",
            "0",
        ])

    programme_table = Table(
        programme_data,
        colWidths=[35 * mm, 120 * mm, 45 * mm, 25 * mm],
    )

    programme_table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ALIGN", (3, 1), (3, -1), "RIGHT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ])
    )

    story.append(programme_table)

    # ==========================================================
    # AWARD
    # ==========================================================

    story.append(
        Paragraph(
            "Enrollment by Award",
            heading_style,
        )
    )

    award_data = [
        [
            "Award",
            "Students",
        ]
    ]

    for award in summary["awards"]:
        award_data.append([
            str(award["programme__award"]).replace("_", " ").title(),
            str(award["total"]),
        ])

    award_table = Table(
        award_data,
        colWidths=[70 * mm, 30 * mm],
    )

    award_table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ALIGN", (1, 1), (1, -1), "RIGHT"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ])
    )

    story.append(award_table)

    # ==========================================================
    # GENDER
    # ==========================================================

    story.append(
        Paragraph(
            "Gender Distribution",
            heading_style,
        )
    )

    gender_data = [
        [
            "Gender",
            "Students",
        ]
    ]

    for item in summary["gender"]:
        gender_data.append([
            item["gender"] or "Not specified",
            str(item["total"]),
        ])

    gender_table = Table(
        gender_data,
        colWidths=[70 * mm, 30 * mm],
    )

    gender_table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ALIGN", (1, 1), (1, -1), "RIGHT"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ])
    )

    story.append(gender_table)

    # ==========================================================
    # STATUS
    # ==========================================================

    story.append(
        Paragraph(
            "Student Status",
            heading_style,
        )
    )

    status_data = [
        [
            "Status",
            "Students",
        ]
    ]

    for item in summary["status"]:
        status_data.append([
            str(item["status"]).replace("_", " ").title(),
            str(item["total"]),
        ])

    status_table = Table(
        status_data,
        colWidths=[70 * mm, 30 * mm],
    )

    status_table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ALIGN", (1, 1), (1, -1), "RIGHT"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ])
    )

    story.append(status_table)

    story.append(Spacer(1, 12))

    story.append(
        Paragraph(
            "XORADEX EDUCORE ERP",
            normal_style,
        )
    )

    document.build(story)

    buffer.seek(0)

    return buffer