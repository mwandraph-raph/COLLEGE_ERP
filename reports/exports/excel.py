"""
XORADEX EDUCORE ERP
EXCEL REPORT EXPORT ENGINE

Provides a common Excel workbook foundation for all reports.
Individual report services will supply the actual data.
"""

from io import BytesIO
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# COMMON WORKBOOK
# ============================================================

def create_workbook(
    title,
    subtitle=None,
):
    """
    Create a standard XORADEX EDUCORE workbook.
    """

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Summary"

    worksheet["A1"] = "XORADEX EDUCORE ERP"
    worksheet["A1"].font = Font(
        bold=True,
        size=16,
    )

    worksheet["A2"] = title
    worksheet["A2"].font = Font(
        bold=True,
        size=13,
    )

    if subtitle:
        worksheet["A3"] = subtitle

    worksheet["A1"].alignment = Alignment(
        vertical="center"
    )

    worksheet["A2"].alignment = Alignment(
        vertical="center"
    )

    return workbook


# ============================================================
# COMMON DETAIL SHEET
# ============================================================

def add_detail_sheet(
    workbook,
    title,
    headers,
    rows,
):
    """
    Add an analysis-friendly detailed data sheet.
    """

    worksheet = workbook.create_sheet(
        title=title[:31]
    )

    for column_number, header in enumerate(
        headers,
        start=1,
    ):

        cell = worksheet.cell(
            row=1,
            column=column_number,
            value=header,
        )

        cell.font = Font(
            bold=True,
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for row_number, row in enumerate(
        rows,
        start=2,
    ):

        for column_number, value in enumerate(
            row,
            start=1,
        ):

            worksheet.cell(
                row=row_number,
                column=column_number,
                value=value,
            )

    worksheet.freeze_panes = "A2"

    if headers:
        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )

    for column_number in range(
        1,
        len(headers) + 1,
    ):

        column_letter = get_column_letter(
            column_number
        )

        maximum_length = len(
            str(
                headers[
                    column_number - 1
                ]
            )
        )

        for cell in worksheet[
            column_letter
        ]:

            if cell.value is not None:

                maximum_length = max(
                    maximum_length,
                    len(str(cell.value)),
                )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            maximum_length + 2,
            40,
        )

    return worksheet


# ============================================================
# EXCEL HTTP RESPONSE
# ============================================================

def workbook_response(
    workbook,
    filename,
):
    """
    Convert workbook into an HTTP response.

    The attachment header forces the Excel file
    to download immediately.
    """

    from django.http import HttpResponse

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )

    response[
        "Content-Disposition"
    ] = (
        f'attachment; filename="{filename}"'
    )

    return response


# ============================================================
# FINANCIAL TRANSACTIONS EXCEL
# ============================================================

def create_financial_transactions_excel(report):
    """
    Create the Financial Transactions Excel workbook.

    The report data is supplied by the financial transaction
    service. This function is responsible only for presentation
    and workbook generation.
    """

    workbook = create_workbook(
        "FINANCIAL TRANSACTIONS & PERFORMANCE REPORT"
    )

    # --------------------------------------------------------
    # EXECUTIVE SUMMARY
    # --------------------------------------------------------

    sheet = workbook.active
    sheet.title = "Executive Summary"

    current = report["current"]
    previous = report["previous"]

    sheet.append(
        [
            "XORADEX EDUCORE",
        ]
    )

    sheet.append(
        [
            "FINANCIAL TRANSACTIONS & PERFORMANCE REPORT",
        ]
    )

    sheet.append([])

    sheet.append(
        [
            "Current Semester",
            (
                f"{current['semester'].academic_year.year_name} - "
                f"{current['semester'].semester_name}"
                if current["semester"]
                else ""
            ),
        ]
    )

    sheet.append(
        [
            "Previous Semester",
            (
                f"{previous['semester'].academic_year.year_name} - "
                f"{previous['semester'].semester_name}"
                if previous["semester"]
                else ""
            ),
        ]
    )

    sheet.append([])

    sheet.append(
        [
            "Indicator",
            "Current",
            "Previous",
            "Difference",
            "% Change",
        ]
    )

    for cell in sheet[7]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    for row in report["comparison"]:

        sheet.append(
            [
                row["label"],

                (
                    float(
                        row["current"]
                        or Decimal("0.00")
                    )
                    if row["key"] != "transaction_count"
                    else row["current"]
                ),

                (
                    float(
                        row["previous"]
                        or Decimal("0.00")
                    )
                    if row["key"] != "transaction_count"
                    else row["previous"]
                ),

                (
                    float(
                        row["difference"]
                        or Decimal("0.00")
                    )
                    if row["key"] != "transaction_count"
                    else row["difference"]
                ),

                float(
                    row["percentage"]
                    or Decimal("0.00")
                ),
            ]
        )

    sheet.append([])

    sheet.append(
        [
            "Current Collection Rate",
            float(
                current["collection_rate"]
                or Decimal("0.00")
            ),
        ]
    )

    sheet.append(
        [
            "Previous Collection Rate",
            float(
                previous["collection_rate"]
                or Decimal("0.00")
            ),
        ]
    )

    sheet.append(
        [
            "Current Outstanding Rate",
            float(
                current["outstanding_rate"]
                or Decimal("0.00")
            ),
        ]
    )

    # --------------------------------------------------------
    # HISTORICAL SEMESTER PERFORMANCE
    # --------------------------------------------------------

    historical_rows = []

    for row in report["historical"]:

        historical_rows.append(
            [
                row["semester"].academic_year.year_name,
                row["semester"].semester_name,
                float(
                    row["expected_revenue"]
                    or Decimal("0.00")
                ),
                float(
                    row["collections"]
                    or Decimal("0.00")
                ),
                float(
                    row["outstanding"]
                    or Decimal("0.00")
                ),
                row["transaction_count"],
                float(
                    row["average_transaction"]
                    or Decimal("0.00")
                ),
                float(
                    row["collection_rate"]
                    or Decimal("0.00")
                ),
                float(
                    row["outstanding_rate"]
                    or Decimal("0.00")
                ),
            ]
        )

    add_detail_sheet(
        workbook,
        "Semester Performance",
        [
            "Academic Year",
            "Semester",
            "Expected Revenue",
            "Collections",
            "Outstanding",
            "Transactions",
            "Average Transaction",
            "Collection Rate",
            "Outstanding Rate",
        ],
        historical_rows,
    )

    # --------------------------------------------------------
    # DEPARTMENT PERFORMANCE
    # --------------------------------------------------------

    department_rows = []

    for row in report["departments"]:

        department_rows.append(
            [
                row["code"],
                row["name"],
                float(
                    row["expected_revenue"]
                    or Decimal("0.00")
                ),
                float(
                    row["collections"]
                    or Decimal("0.00")
                ),
                float(
                    row["outstanding"]
                    or Decimal("0.00")
                ),
                float(
                    row["collection_rate"]
                    or Decimal("0.00")
                ),
            ]
        )

    add_detail_sheet(
        workbook,
        "Department Performance",
        [
            "Code",
            "Department",
            "Expected Revenue",
            "Collections",
            "Outstanding",
            "Collection Rate",
        ],
        department_rows,
    )

    # --------------------------------------------------------
    # DETAILED POSTED TRANSACTIONS
    # --------------------------------------------------------

    transaction_rows = []

    for payment in report["transactions"]:

        student = payment.invoice.student

        department = (
            student.programme.course.department
            if student.programme
            and student.programme.course
            else None
        )

        received_by = ""

        if payment.received_by:

            received_by = (
                payment.received_by.get_full_name()
                or payment.received_by.username
            )

        transaction_rows.append(
            [
                payment.payment_number,
                payment.invoice.invoice_number,
                student.admission_no,

                (
                    f"{student.first_name} "
                    f"{student.middle_name} "
                    f"{student.last_name}"
                ).strip(),

                (
                    payment.invoice.enrollment
                    .academic_year.year_name
                    if payment.invoice.enrollment
                    else ""
                ),

                (
                    payment.invoice.enrollment
                    .semester.semester_name
                    if payment.invoice.enrollment
                    else ""
                ),

                department.name if department else "",

                (
                    student.programme.name
                    if student.programme
                    else ""
                ),

                payment.payment_date,

                float(
                    payment.amount
                    or Decimal("0.00")
                ),

                payment.payment_method,

                payment.reference_number or "",

                received_by,
            ]
        )

    add_detail_sheet(
        workbook,
        "Posted Transactions",
        [
            "Payment Number",
            "Invoice Number",
            "Admission Number",
            "Student",
            "Academic Year",
            "Semester",
            "Department",
            "Programme",
            "Payment Date",
            "Amount",
            "Payment Method",
            "Reference Number",
            "Received By",
        ],
        transaction_rows,
    )

    # --------------------------------------------------------
    # FORMAT EXECUTIVE SUMMARY
    # --------------------------------------------------------

    for column_cells in sheet.columns:

        maximum_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            if cell.value is not None:

                maximum_length = max(
                    maximum_length,
                    len(str(cell.value)),
                )

        sheet.column_dimensions[
            column_letter
        ].width = min(
            max(maximum_length + 2, 12),
            45,
        )

    # --------------------------------------------------------
    # RETURN WORKBOOK
    # --------------------------------------------------------

    return workbook