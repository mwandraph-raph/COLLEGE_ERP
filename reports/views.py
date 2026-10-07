from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse, FileResponse
from django.shortcuts import (
    render,
    get_object_or_404,
)
from django.contrib.auth import get_user_model
from system.models import ActivityLog
import io

from openpyxl import Workbook
from openpyxl.styles import Font

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from .registry import get_reports_by_category
from .services.system import (
    get_user_accounts_report,
    get_audit_trail_report,
)
from .services.students import (
    get_student_register,
    student_register_rows,
)


from .services.academic import (
    get_examination_performance_report,
    get_missing_marks_report,
)

from .services.staff import (
    get_staff_register_report,
    get_staff_roles,
)

from .services.enrollment import (
    get_enrollment_summary,
)

from .services.student_status import (
    get_student_status_summary,
)

from .services.financial_position import (
    get_financial_position_summary,
)

from .services.financial_transactions import (
    get_financial_transactions_report,
    get_current_semester,
    get_management_analysis,
    calculate_semester_metrics,
)

from .services.financial_transactions import (
    get_all_semesters,
    get_outstanding_fees_report,
)

from .exports.excel import (
    create_workbook,
    add_detail_sheet,
    workbook_response,
    create_financial_transactions_excel,
)

from .exports.pdf import (
    create_student_register_pdf,
    pdf_response,
    create_financial_transactions_pdf,
)

from .exports.enrollment import (
    create_enrollment_excel,
    create_enrollment_pdf,
)


from students.models import (
    AcademicYear,
    Semester,
    Department,
    Programme,
    Student,
    Result,
)

from accounts.models import User


# ============================================================
# REPORT ROLE
# ============================================================

def get_report_role(user):

    if user.is_superuser:
        return "ADMIN"

    group_to_role = {
        "Administrator": "ADMIN",
        "Principal": "PRINCIPAL",
        "HOD": "HOD",
        "Registrar": "REGISTRAR",
        "Finance Officer": "FINANCE_OFFICER",
        "Exam Officer": "EXAMINATIONS_OFFICER",
        "Graduation": "GRADUATION_OFFICER",
    }

    for group_name, role in group_to_role.items():

        if user.groups.filter(
            name=group_name
        ).exists():

            return role

    return ""


# ============================================================
# REPORT CENTER
# ============================================================

@login_required
def report_center(request):

    role = get_report_role(
        request.user
    )

    report_categories = get_reports_by_category(
        role
    )

    return render(
        request,
        "reports/center.html",
        {
            "report_role": role,
            "report_categories": report_categories,
        },
    )


# ============================================================
# STUDENT STATUS REPORT
# ============================================================

@login_required
def student_status_report(request):
    """
    Student Status Report.

    Student numbers grouped by current status.

    HOD users see only students belonging
    to their department.
    """

    summary = get_student_status_summary(
        request.user
    )

    return render(
        request,
        "reports/student_status.html",
        {
            **summary,
            "report_role": get_report_role(
                request.user
            ),
        },
    )


# ============================================================
# STUDENT REGISTER HELPERS
# ============================================================

def get_student_register_queryset(request):

    department = (
        request.GET.get("department")
        or None
    )

    course = (
        request.GET.get("course")
        or None
    )

    programme = (
        request.GET.get("programme")
        or None
    )

    status = (
        request.GET.get("status")
        or None
    )

    gender = (
        request.GET.get("gender")
        or None
    )

    search = (
        request.GET.get("search")
        or None
    )

    admission_date_from = (
        request.GET.get(
            "admission_date_from"
        )
        or None
    )

    admission_date_to = (
        request.GET.get(
            "admission_date_to"
        )
        or None
    )

    students = get_student_register(
        department=department,
        course=course,
        programme=programme,
        status=status,
        gender=gender,
        search=search,
        admission_date_from=admission_date_from,
        admission_date_to=admission_date_to,
    )

    if request.user.groups.filter(
        name="HOD"
    ).exists():

        students = students.filter(
            programme__course__department__hod=request.user
        )

    return students


def get_student_register_dropdowns(request):

    from students.models import Student

    students = (
        Student.objects
        .select_related(
            "programme",
            "programme__course",
            "programme__course__department",
        )
    )

    if request.user.groups.filter(
        name="HOD"
    ).exists():

        students = students.filter(
            programme__course__department__hod=request.user
        )

    departments = (
        students
        .values(
            "programme__course__department_id",
            "programme__course__department__name",
        )
        .distinct()
        .order_by(
            "programme__course__department__name"
        )
    )

    courses = (
        students
        .values(
            "programme__course_id",
            "programme__course__name",
        )
        .distinct()
        .order_by(
            "programme__course__name"
        )
    )

    programmes = (
        students
        .values(
            "programme_id",
            "programme__name",
        )
        .distinct()
        .order_by(
            "programme__name"
        )
    )

    status_choices = Student._meta.get_field(
        "status"
    ).choices

    gender_choices = Student._meta.get_field(
        "gender"
    ).choices

    return {
        "departments": departments,
        "courses": courses,
        "programmes": programmes,
        "status_choices": status_choices,
        "gender_choices": gender_choices,
    }


# ============================================================
# STUDENT REGISTER
# ============================================================

@login_required
def student_register_report(request):

    students = get_student_register_queryset(
        request
    )

    rows = student_register_rows(
        students
    )

    dropdowns = get_student_register_dropdowns(
        request
    )

    return render(
        request,
        "reports/student_register.html",
        {
            "rows": rows,
            **dropdowns,

            "selected_department":
                request.GET.get(
                    "department",
                    "",
                ),

            "selected_course":
                request.GET.get(
                    "course",
                    "",
                ),

            "selected_programme":
                request.GET.get(
                    "programme",
                    "",
                ),

            "selected_status":
                request.GET.get(
                    "status",
                    "",
                ),

            "selected_gender":
                request.GET.get(
                    "gender",
                    "",
                ),

            "search":
                request.GET.get(
                    "search",
                    "",
                ),

            "admission_date_from":
                request.GET.get(
                    "admission_date_from",
                    "",
                ),

            "admission_date_to":
                request.GET.get(
                    "admission_date_to",
                    "",
                ),

            "report_role":
                get_report_role(
                    request.user
                ),
        },
    )


# ============================================================
# STUDENT REGISTER AJAX SEARCH
# ============================================================

@login_required
def student_register_ajax_search(request):

    students = get_student_register_queryset(
        request
    )

    rows = student_register_rows(
        students
    )

    return JsonResponse(
        {
            "success": True,
            "count": len(rows),
            "rows": rows,
        }
    )


# ============================================================
# STUDENT REGISTER EXCEL
# ============================================================

@login_required
def student_register_excel(request):

    students = get_student_register_queryset(
        request
    )

    rows = student_register_rows(
        students
    )

    workbook = create_workbook(
        title="Student Master Register",
        subtitle=f"Total Students: {len(rows)}",
    )

    headers = [
        "Admission No.",
        "Student Name",
        "Gender",
        "Date of Birth",
        "ID Number",
        "Phone",
        "Email",
        "Department",
        "Course",
        "Programme",
        "Admission Date",
        "Status",
    ]

    excel_rows = []

    for row in rows:

        excel_rows.append(
            [
                row["admission_no"],
                row["name"],
                row["gender"],
                row["date_of_birth"],
                row["id_number"],
                row["phone"],
                row["email"],
                row["department"],
                row["course"],
                row["programme"],
                row["admission_date"],
                row["status"],
            ]
        )

    add_detail_sheet(
        workbook,
        title="Student Register",
        headers=headers,
        rows=excel_rows,
    )

    return workbook_response(
        workbook,
        "XORADEX_Student_Master_Register.xlsx",
    )


# ============================================================
# STUDENT REGISTER PDF
# ============================================================

@login_required
def student_register_pdf(request):

    students = get_student_register_queryset(
        request
    )

    rows = student_register_rows(
        students
    )

    buffer = create_student_register_pdf(
        rows
    )

    return pdf_response(
        buffer,
        "XORADEX_Student_Master_Register.pdf",
    )


# ============================================================
# FINANCIAL TRANSACTIONS
# ============================================================

@login_required
def financial_transactions(request):

    # ========================================================
    # ALL SEMESTERS FOR REPORTING
    # Do NOT filter by is_active
    # ========================================================

    semesters = (
        Semester.objects
        .select_related(
            "academic_year"
        )
        .all()
        .order_by(
            "-academic_year__year_name",
            "-start_date",
            "-id",
        )
    )

    # ========================================================
    # SELECTED REPORTING SEMESTER
    # ========================================================

    semester_id = request.GET.get(
        "semester"
    )

    if semester_id:

        try:

            current_semester = semesters.get(
                pk=int(semester_id)
            )

        except (
            Semester.DoesNotExist,
            ValueError,
            TypeError,
        ):

            current_semester = get_current_semester(
                request.user
            )

    else:

        current_semester = get_current_semester(
            request.user
        )

    # ========================================================
    # FILTERS
    # ========================================================

    department_id = (
        request.GET.get(
            "department"
        )
        or None
    )

    programme_id = (
        request.GET.get(
            "programme"
        )
        or None
    )

    # ========================================================
    # AJAX SEMESTER COMPARISON
    # ========================================================

    if (
        request.headers.get(
            "X-Requested-With"
        )
        == "XMLHttpRequest"
        and request.GET.get(
            "ajax"
        )
        == "semester_comparison"
    ):

        trend_semester_ids = request.GET.getlist(
            "trend_semester"
        )

        valid_semester_ids = []

        for value in trend_semester_ids:

            try:

                semester_pk = int(
                    value
                )

            except (
                ValueError,
                TypeError,
            ):

                continue

            if semester_pk not in valid_semester_ids:

                valid_semester_ids.append(
                    semester_pk
                )

        if len(valid_semester_ids) < 2:

            return JsonResponse(
                {
                    "historical": [],
                    "error": (
                        "Please select at least two "
                        "semesters to compare."
                    ),
                },
                status=400,
            )

        selected_semesters = list(
            semesters.filter(
                pk__in=valid_semester_ids
            )
        )

        semester_map = {
            semester.id: semester
            for semester in selected_semesters
        }

        ordered_semesters = [
            semester_map[semester_pk]
            for semester_pk in valid_semester_ids
            if semester_pk in semester_map
        ]

        historical = []

        for semester in ordered_semesters:

            metrics = calculate_semester_metrics(
                request.user,
                semester,
                department_id=department_id,
                programme_id=programme_id,
            )

            academic_year = getattr(
                semester,
                "academic_year",
                None,
            )

            academic_year_name = (
                getattr(
                    academic_year,
                    "year_name",
                    None,
                )
                or "-"
            )

            semester_name = (
                getattr(
                    semester,
                    "semester_name",
                    None,
                )
                or "-"
            )

            historical.append(
                {
                    "academic_year":
                        academic_year_name,

                    "semester":
                        semester_name,

                    "semester_id":
                        semester.id,

                    "gross_invoiced":
                        float(
                            metrics.get(
                                "gross_invoiced",
                                0,
                            )
                            or 0
                        ),

                    "cash_collections":
                        float(
                            metrics.get(
                                "cash_collections",
                                0,
                            )
                            or 0
                        ),

                    "credits_applied":
                        float(
                            metrics.get(
                                "credits_applied",
                                0,
                            )
                            or 0
                        ),

                    "outstanding":
                        float(
                            metrics.get(
                                "outstanding",
                                0,
                            )
                            or 0
                        ),

                    "overpayment_credit":
                        float(
                            metrics.get(
                                "overpayment_credit",
                                0,
                            )
                            or 0
                        ),

                    "settlement_rate":
                        float(
                            metrics.get(
                                "settlement_rate",
                                0,
                            )
                            or 0
                        ),

                    "cash_collection_rate":
                        float(
                            metrics.get(
                                "cash_collection_rate",
                                0,
                            )
                            or 0
                        ),

                    "transaction_count":
                        int(
                            metrics.get(
                                "transaction_count",
                                0,
                            )
                            or 0
                        ),

                    "average_transaction":
                        float(
                            metrics.get(
                                "average_transaction",
                                0,
                            )
                            or 0
                        ),
                }
            )

        return JsonResponse(
            {
                "historical": historical,
            }
        )

    # ========================================================
    # NORMAL REPORT
    # ========================================================

    report = get_financial_transactions_report(
        request.user,
        current_semester=current_semester,
        department_id=department_id,
        programme_id=programme_id,
    )

    # ========================================================
    # MANAGEMENT ANALYSIS
    # ========================================================

    management_analysis = get_management_analysis(
        report["current"],
        report["previous"],
    )

    # ========================================================
    # DEPARTMENTS
    # ========================================================

    departments = (
        Department.objects
        .filter(
            is_active=True
        )
        .order_by(
            "name"
        )
    )

    # ========================================================
    # PROGRAMMES
    # ========================================================

    programmes = (
        Programme.objects
        .filter(
            is_active=True
        )
        .select_related(
            "course",
            "course__department",
        )
        .order_by(
            "name"
        )
    )

    # ========================================================
    # HOD RESTRICTION
    # ========================================================

    if request.user.groups.filter(
        name="HOD"
    ).exists():

        departments = departments.filter(
            hod=request.user
        )

        programmes = programmes.filter(
            course__department__hod=request.user
        )

    # ========================================================
    # RENDER
    # ========================================================

    return render(
        request,
        "reports/financial_transactions.html",
        {
            "report":
                report,

            "current":
                report["current"],

            "previous":
                report["previous"],

            "comparison":
                report["comparison"],

            "historical":
                report["historical"],

            "department_performance":
                report["department_performance"],

            "transactions":
                report["transactions"],

            "current_semester":
                report["current_semester"],

            "previous_semester":
                report["previous_semester"],

            "management_analysis":
                management_analysis,

            "departments":
                departments,

            "programmes":
                programmes,

            "semesters":
                semesters,

            "selected_department":
                department_id,

            "selected_programme":
                programme_id,

            "selected_semester":
                current_semester.id
                if current_semester
                else None,

            "report_role":
                get_report_role(
                    request.user
                ),
        },
    )


# ============================================================
# ENROLLMENT SUMMARY
# ============================================================

@login_required
def enrollment_summary(request):
    """
    Enrollment Summary report.

    HOD users see only students belonging
    to their department.
    """

    summary = get_enrollment_summary(
        request.user
    )

    return render(
        request,
        "reports/enrollment_summary.html",
        {
            **summary,
            "report_role":
                get_report_role(
                    request.user
                ),
        },
    )


# ============================================================
# ENROLLMENT SUMMARY EXCEL
# ============================================================

@login_required
def enrollment_summary_excel(request):
    """
    Export Enrollment Summary to Excel.
    """

    summary = get_enrollment_summary(
        request.user
    )

    workbook = create_enrollment_excel(
        summary
    )

    response = HttpResponse(
        workbook.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )

    response[
        "Content-Disposition"
    ] = (
        'attachment; '
        'filename="XORADEX_Enrollment_Summary.xlsx"'
    )

    return response


# ============================================================
# ENROLLMENT SUMMARY PDF
# ============================================================

@login_required
def enrollment_summary_pdf(request):
    """
    Export Enrollment Summary to PDF.
    """

    summary = get_enrollment_summary(
        request.user
    )

    pdf = create_enrollment_pdf(
        summary
    )

    response = HttpResponse(
        pdf.getvalue(),
        content_type="application/pdf",
    )

    response[
        "Content-Disposition"
    ] = (
        'attachment; '
        'filename="XORADEX_Enrollment_Summary.pdf"'
    )

    return response


# ============================================================
# FINANCIAL POSITION
# ============================================================

@login_required
def financial_position(request):
    """
    Financial Position Report.

    Shows expected revenue, collections, credits
    and outstanding balances.

    HOD users see only financial information
    belonging to their department.
    """

    summary = get_financial_position_summary(
        request.user
    )

    return render(
        request,
        "reports/financial_position.html",
        {
            **summary,
            "report_role":
                get_report_role(
                    request.user
                ),
        },
    )


# ============================================================
# FINANCIAL TRANSACTIONS EXCEL
# ============================================================

@login_required
def financial_transactions_excel(request):

    semesters = (
        Semester.objects
        .select_related(
            "academic_year"
        )
        .all()
        .order_by(
            "-academic_year__year_name",
            "-start_date",
            "-id",
        )
    )

    semester_id = request.GET.get(
        "semester"
    )

    if semester_id:

        try:

            current_semester = semesters.get(
                pk=int(semester_id)
            )

        except (
            Semester.DoesNotExist,
            ValueError,
            TypeError,
        ):

            current_semester = get_current_semester(
                request.user
            )

    else:

        current_semester = get_current_semester(
            request.user
        )

    department_id = (
        request.GET.get(
            "department"
        )
        or None
    )

    programme_id = (
        request.GET.get(
            "programme"
        )
        or None
    )

    report = get_financial_transactions_report(
        request.user,
        current_semester=current_semester,
        department_id=department_id,
        programme_id=programme_id,
    )

    current = report["current"]
    previous = report["previous"]
    historical = report["historical"]

    department_performance = report[
        "department_performance"
    ]

    transactions = report[
        "transactions"
    ]

    workbook = Workbook()

    # ========================================================
    # SUMMARY SHEET
    # ========================================================

    worksheet = workbook.active

    worksheet.title = "Financial Summary"

    worksheet["A1"] = "XORADEX EDUCORE"

    worksheet["A1"].font = Font(
        bold=True,
        size=16,
    )

    worksheet["A2"] = (
        "Financial Transactions Report"
    )

    worksheet["A2"].font = Font(
        bold=True,
        size=13,
    )

    academic_year = getattr(
        getattr(
            current_semester,
            "academic_year",
            None,
        ),
        "year_name",
        "",
    )

    semester_name = getattr(
        current_semester,
        "semester_name",
        "",
    )

    worksheet["A4"] = "Academic Year"
    worksheet["B4"] = academic_year or "—"

    worksheet["A5"] = "Semester"
    worksheet["B5"] = semester_name or "—"

    worksheet["A7"] = (
        "Financial Performance"
    )

    worksheet["A7"].font = Font(
        bold=True
    )

    summary_rows = [
        (
            "Gross Invoiced",
            current.get(
                "gross_invoiced",
                0,
            ),
        ),
        (
            "Cash Collections",
            current.get(
                "cash_collections",
                0,
            ),
        ),
        (
            "Credits Applied",
            current.get(
                "credits_applied",
                0,
            ),
        ),
        (
            "Outstanding",
            current.get(
                "outstanding",
                0,
            ),
        ),
        (
            "Credit / Overpayment Position",
            current.get(
                "overpayment_credit",
                0,
            ),
        ),
        (
            "Settlement Rate",
            current.get(
                "settlement_rate",
                0,
            ),
        ),
        (
            "Cash Collection Rate",
            current.get(
                "cash_collection_rate",
                0,
            ),
        ),
        (
            "Transactions",
            current.get(
                "transaction_count",
                0,
            ),
        ),
        (
            "Average Transaction",
            current.get(
                "average_transaction",
                0,
            ),
        ),
    ]

    row_number = 8

    for label, value in summary_rows:

        worksheet.cell(
            row=row_number,
            column=1,
            value=label,
        )

        worksheet.cell(
            row=row_number,
            column=2,
            value=float(
                value or 0
            ),
        )

        row_number += 1

    # ========================================================
    # HISTORICAL TREND
    # ========================================================

    history_sheet = workbook.create_sheet(
        "Financial Trend"
    )

    history_headers = [
        "Academic Year",
        "Semester",
        "Gross Invoiced",
        "Cash Collections",
        "Credits Applied",
        "Outstanding",
        "Settlement Rate",
        "Cash Collection Rate",
        "Transactions",
    ]

    history_sheet.append(
        history_headers
    )

    for cell in history_sheet[1]:

        cell.font = Font(
            bold=True
        )

    for item in historical:

        history_sheet.append(
            [
                item.get(
                    "academic_year",
                    "—",
                ),

                item.get(
                    "semester",
                    "—",
                ),

                float(
                    item.get(
                        "gross_invoiced",
                        0,
                    )
                    or 0
                ),

                float(
                    item.get(
                        "cash_collections",
                        0,
                    )
                    or 0
                ),

                float(
                    item.get(
                        "credits_applied",
                        0,
                    )
                    or 0
                ),

                float(
                    item.get(
                        "outstanding",
                        0,
                    )
                    or 0
                ),

                float(
                    item.get(
                        "settlement_rate",
                        0,
                    )
                    or 0
                ),

                float(
                    item.get(
                        "cash_collection_rate",
                        0,
                    )
                    or 0
                ),

                int(
                    item.get(
                        "transaction_count",
                        0,
                    )
                    or 0
                ),
            ]
        )

    # ========================================================
    # DEPARTMENT PERFORMANCE
    # ========================================================

    department_sheet = workbook.create_sheet(
        "Department Performance"
    )

    department_headers = [
        "Department",
        "Code",
        "Gross Invoiced",
        "Cash Collections",
        "Credits Applied",
        "Outstanding",
        "Settlement Rate",
        "Cash Collection Rate",
    ]

    department_sheet.append(
        department_headers
    )

    for cell in department_sheet[1]:

        cell.font = Font(
            bold=True
        )

    for item in department_performance:

        department_sheet.append(
            [
                item.get(
                    "department_name",
                    "—",
                ),

                item.get(
                    "department_code",
                    "—",
                ),

                float(
                    item.get(
                        "gross_invoiced",
                        0,
                    )
                    or 0
                ),

                float(
                    item.get(
                        "cash_collections",
                        0,
                    )
                    or 0
                ),

                float(
                    item.get(
                        "credits_applied",
                        0,
                    )
                    or 0
                ),

                float(
                    item.get(
                        "outstanding",
                        0,
                    )
                    or 0
                ),

                float(
                    item.get(
                        "settlement_rate",
                        0,
                    )
                    or 0
                ),

                float(
                    item.get(
                        "cash_collection_rate",
                        0,
                    )
                    or 0
                ),
            ]
        )

    # ========================================================
    # POSTED TRANSACTIONS
    # ========================================================

    transaction_sheet = workbook.create_sheet(
        "Posted Transactions"
    )

    transaction_headers = [
        "Payment Number",
        "Invoice Number",
        "Student",
        "Payment Date",
        "Amount",
        "Payment Method",
        "Reference Number",
        "Received By",
    ]

    transaction_sheet.append(
        transaction_headers
    )

    for cell in transaction_sheet[1]:

        cell.font = Font(
            bold=True
        )

    for payment in transactions:

        student = getattr(
            payment.invoice,
            "student",
            None,
        )

        student_name = ""

        if student:

            student_name = (
                f"{getattr(student, 'first_name', '')} "
                f"{getattr(student, 'middle_name', '')} "
                f"{getattr(student, 'last_name', '')}"
            ).strip()

        received_by = getattr(
            payment.received_by,
            "username",
            "",
        )

        transaction_sheet.append(
            [
                getattr(
                    payment,
                    "payment_number",
                    "",
                ),

                getattr(
                    payment.invoice,
                    "invoice_number",
                    "",
                ),

                student_name,

                payment.payment_date.strftime(
                    "%Y-%m-%d"
                )
                if payment.payment_date
                else "",

                float(
                    payment.amount
                    or 0
                ),

                getattr(
                    payment,
                    "payment_method",
                    "",
                ),

                getattr(
                    payment,
                    "reference_number",
                    "",
                ),

                received_by,
            ]
        )

    # ========================================================
    # COLUMN WIDTHS
    # ========================================================

    for sheet in workbook.worksheets:

        for column_cells in sheet.columns:

            maximum_length = 0

            column_letter = (
                column_cells[0].column_letter
            )

            for cell in column_cells:

                try:

                    value_length = len(
                        str(
                            cell.value
                        )
                    )

                    maximum_length = max(
                        maximum_length,
                        value_length,
                    )

                except Exception:

                    pass

            sheet.column_dimensions[
                column_letter
            ].width = min(
                maximum_length + 2,
                35,
            )

    # ========================================================
    # RESPONSE
    # ========================================================

    output = io.BytesIO()

    workbook.save(
        output
    )

    output.seek(0)

    filename = (
        "financial_transactions_"
        f"{academic_year}_"
        f"{semester_name}.xlsx"
    )

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
    )

    response[
        "Content-Disposition"
    ] = (
        f'attachment; filename="{filename}"'
    )

    return response


# ============================================================
# FINANCIAL TRANSACTIONS PDF
# ============================================================

@login_required
def financial_transactions_pdf(request):

    semesters = (
        Semester.objects
        .select_related(
            "academic_year"
        )
        .all()
        .order_by(
            "-academic_year__year_name",
            "-start_date",
            "-id",
        )
    )

    semester_id = request.GET.get(
        "semester"
    )

    if semester_id:

        try:

            current_semester = semesters.get(
                pk=int(semester_id)
            )

        except (
            Semester.DoesNotExist,
            ValueError,
            TypeError,
        ):

            current_semester = get_current_semester(
                request.user
            )

    else:

        current_semester = get_current_semester(
            request.user
        )

    department_id = (
        request.GET.get(
            "department"
        )
        or None
    )

    programme_id = (
        request.GET.get(
            "programme"
        )
        or None
    )

    report = get_financial_transactions_report(
        request.user,
        current_semester=current_semester,
        department_id=department_id,
        programme_id=programme_id,
    )

    current = report["current"]
    previous = report["previous"]
    historical = report["historical"]

    department_performance = report[
        "department_performance"
    ]

    transactions = report[
        "transactions"
    ]

    management_analysis = get_management_analysis(
        current,
        previous,
    )

    buffer = io.BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
    )

    styles = getSampleStyleSheet()

    title_style = styles["Title"]

    title_style.alignment = TA_CENTER

    heading_style = styles["Heading2"]
    normal_style = styles["Normal"]

    story = []

    # ========================================================
    # TITLE
    # ========================================================

    story.append(
        Paragraph(
            "XORADEX EDUCORE",
            title_style,
        )
    )

    story.append(
        Paragraph(
            "Financial Transactions Report",
            heading_style,
        )
    )

    story.append(
        Spacer(
            1,
            5 * mm,
        )
    )

    # ========================================================
    # SEMESTER INFORMATION
    # ========================================================

    academic_year = getattr(
        getattr(
            current_semester,
            "academic_year",
            None,
        ),
        "year_name",
        "—",
    ) or "—"

    semester_name = getattr(
        current_semester,
        "semester_name",
        "—",
    ) or "—"

    story.append(
        Paragraph(
            f"<b>Academic Year:</b> "
            f"{academic_year}",
            normal_style,
        )
    )

    story.append(
        Paragraph(
            f"<b>Semester:</b> "
            f"{semester_name}",
            normal_style,
        )
    )

    if department_id:

        story.append(
            Paragraph(
                f"<b>Department Filter:</b> "
                f"{department_id}",
                normal_style,
            )
        )

    if programme_id:

        story.append(
            Paragraph(
                f"<b>Programme Filter:</b> "
                f"{programme_id}",
                normal_style,
            )
        )

    story.append(
        Spacer(
            1,
            5 * mm,
        )
    )

    # ========================================================
    # FINANCIAL SUMMARY
    # ========================================================

    story.append(
        Paragraph(
            "Financial Performance",
            heading_style,
        )
    )

    summary_data = [
        [
            "Gross Invoiced",
            "Cash Collections",
            "Credits Applied",
            "Outstanding",
            "Overpayment",
            "Settlement",
            "Cash Rate",
            "Transactions",
        ],
        [
            f"KSh {float(current.get('gross_invoiced', 0) or 0):,.2f}",
            f"KSh {float(current.get('cash_collections', 0) or 0):,.2f}",
            f"KSh {float(current.get('credits_applied', 0) or 0):,.2f}",
            f"KSh {float(current.get('outstanding', 0) or 0):,.2f}",
            f"KSh {float(current.get('overpayment_credit', 0) or 0):,.2f}",
            f"{float(current.get('settlement_rate', 0) or 0):.1f}%",
            f"{float(current.get('cash_collection_rate', 0) or 0):.1f}%",
            str(
                current.get(
                    "transaction_count",
                    0,
                )
                or 0
            ),
        ],
    ]

    summary_table = Table(
        summary_data,
        repeatRows=1,
    )

    summary_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey,
                ),
                (
                    "ALIGN",
                    (1, 0),
                    (-1, -1),
                    "RIGHT",
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
            ]
        )
    )

    story.append(
        summary_table
    )

    story.append(
        Spacer(
            1,
            7 * mm,
        )
    )

    # ========================================================
    # FINANCIAL TREND
    # ========================================================

    story.append(
        Paragraph(
            "Financial Trend",
            heading_style,
        )
    )

    trend_data = [
        [
            "Academic Year",
            "Semester",
            "Gross Invoiced",
            "Cash Collections",
            "Credits Applied",
            "Outstanding",
            "Settlement",
            "Cash Rate",
            "Transactions",
        ]
    ]

    for row in historical:

        trend_data.append(
            [
                row.get(
                    "academic_year",
                    "—",
                ),

                row.get(
                    "semester",
                    "—",
                ),

                f"KSh {float(row.get('gross_invoiced', 0) or 0):,.2f}",

                f"KSh {float(row.get('cash_collections', 0) or 0):,.2f}",

                f"KSh {float(row.get('credits_applied', 0) or 0):,.2f}",

                f"KSh {float(row.get('outstanding', 0) or 0):,.2f}",

                f"{float(row.get('settlement_rate', 0) or 0):.1f}%",

                f"{float(row.get('cash_collection_rate', 0) or 0):.1f}%",

                str(
                    row.get(
                        "transaction_count",
                        0,
                    )
                    or 0
                ),
            ]
        )

    trend_table = Table(
        trend_data,
        repeatRows=1,
    )

    trend_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey,
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
                (
                    "ALIGN",
                    (2, 1),
                    (-1, -1),
                    "RIGHT",
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
            ]
        )
    )

    story.append(
        trend_table
    )

    story.append(
        Spacer(
            1,
            7 * mm,
        )
    )

    # ========================================================
    # DEPARTMENT PERFORMANCE
    # ========================================================

    if department_performance:

        story.append(
            Paragraph(
                "Department Performance",
                heading_style,
            )
        )

        department_data = [
            [
                "Department",
                "Code",
                "Gross Invoiced",
                "Cash Collections",
                "Credits Applied",
                "Outstanding",
                "Settlement",
                "Cash Rate",
            ]
        ]

        for row in department_performance:

            department_data.append(
                [
                    row.get(
                        "department_name",
                        "—",
                    ),

                    row.get(
                        "department_code",
                        "—",
                    ),

                    f"KSh {float(row.get('gross_invoiced', 0) or 0):,.2f}",

                    f"KSh {float(row.get('cash_collections', 0) or 0):,.2f}",

                    f"KSh {float(row.get('credits_applied', 0) or 0):,.2f}",

                    f"KSh {float(row.get('outstanding', 0) or 0):,.2f}",

                    f"{float(row.get('settlement_rate', 0) or 0):.1f}%",

                    f"{float(row.get('cash_collection_rate', 0) or 0):.1f}%",
                ]
            )

        department_table = Table(
            department_data,
            repeatRows=1,
        )

        department_table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.lightgrey,
                    ),
                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "Helvetica-Bold",
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.5,
                        colors.grey,
                    ),
                    (
                        "FONTSIZE",
                        (0, 0),
                        (-1, -1),
                        7,
                    ),
                    (
                        "ALIGN",
                        (2, 1),
                        (-1, -1),
                        "RIGHT",
                    ),
                ]
            )
        )

        story.append(
            department_table
        )

        story.append(
            Spacer(
                1,
                7 * mm,
            )
        )

    # ========================================================
    # POSTED TRANSACTIONS
    # ========================================================

    story.append(
        Paragraph(
            "Posted Transactions",
            heading_style,
        )
    )

    transaction_data = [
        [
            "Payment No.",
            "Invoice No.",
            "Student",
            "Date",
            "Amount",
            "Method",
            "Reference",
        ]
    ]

    for payment in transactions:

        student = getattr(
            payment.invoice,
            "student",
            None,
        )

        student_name = ""

        if student:

            student_name = (
                f"{getattr(student, 'first_name', '')} "
                f"{getattr(student, 'middle_name', '')} "
                f"{getattr(student, 'last_name', '')}"
            ).strip()

        transaction_data.append(
            [
                getattr(
                    payment,
                    "payment_number",
                    "",
                ),

                getattr(
                    payment.invoice,
                    "invoice_number",
                    "",
                ),

                student_name,

                payment.payment_date.strftime(
                    "%Y-%m-%d"
                )
                if payment.payment_date
                else "",

                f"KSh {float(payment.amount or 0):,.2f}",

                str(
                    getattr(
                        payment,
                        "payment_method",
                        "",
                    )
                ),

                getattr(
                    payment,
                    "reference_number",
                    "",
                )
                or "",
            ]
        )

    if len(transaction_data) == 1:

        transaction_data.append(
            [
                "No posted transactions",
                "",
                "",
                "",
                "",
                "",
                "",
            ]
        )

    transaction_table = Table(
        transaction_data,
        repeatRows=1,
    )

    transaction_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey,
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
                (
                    "ALIGN",
                    (4, 1),
                    (4, -1),
                    "RIGHT",
                ),
            ]
        )
    )

    story.append(
        transaction_table
    )

    document.build(
        story
    )

    buffer.seek(0)

    filename = (
        "financial_transactions_"
        f"{academic_year}_"
        f"{semester_name}.pdf"
    )

    return FileResponse(
        buffer,
        as_attachment=True,
        filename=filename,
        content_type="application/pdf",
    )



# ============================================================
# OUTSTANDING FEES
# ============================================================

@login_required
def outstanding_fees(request):

    semester_id = (
        request.GET.get(
            "semester"
        )
        or None
    )

    department_id = (
        request.GET.get(
            "department"
        )
        or None
    )

    programme_id = (
        request.GET.get(
            "programme"
        )
        or None
    )

    search = (
        request.GET.get(
            "search"
        )
        or ""
    ).strip()

    report = get_outstanding_fees_report(
        request.user,
        semester_id=semester_id,
        department_id=department_id,
        programme_id=programme_id,
        search=search,
    )

    semesters = (
        get_all_semesters(
            request.user
        )
        .order_by(
            "-academic_year__year_name",
            "-start_date",
            "-id",
        )
    )

    departments = (
        Department.objects
        .filter(
            is_active=True
        )
        .order_by(
            "name"
        )
    )

    programmes = (
        Programme.objects
        .filter(
            is_active=True
        )
        .select_related(
            "course",
            "course__department",
        )
        .order_by(
            "name"
        )
    )

    # --------------------------------------------------------
    # HOD RESTRICTION
    # --------------------------------------------------------

    if request.user.groups.filter(
        name="HOD"
    ).exists():

        departments = departments.filter(
            hod=request.user
        )

        programmes = programmes.filter(
            course__department__hod=request.user
        )

    # --------------------------------------------------------
    # RENDER
    # --------------------------------------------------------

    return render(
        request,
        "reports/outstanding_fees.html",
        {
            "report":
                report,

            "rows":
                report["rows"],

            "summary":
                report["summary"],

            "semesters":
                semesters,

            "departments":
                departments,

            "programmes":
                programmes,

            "selected_semester":
                report["semester_id"],

            "selected_department":
                department_id,

            "selected_programme":
                programme_id,

            "search":
                search,

            "report_role":
                get_report_role(
                    request.user
                ),
        },
    )


# ============================================================
# OUTSTANDING FEES EXCEL
# ============================================================

@login_required
def outstanding_fees_excel(request):

    semester_id = (
        request.GET.get(
            "semester"
        )
        or None
    )

    department_id = (
        request.GET.get(
            "department"
        )
        or None
    )

    programme_id = (
        request.GET.get(
            "programme"
        )
        or None
    )

    search = (
        request.GET.get(
            "search"
        )
        or ""
    ).strip()

    report = get_outstanding_fees_report(
        request.user,
        semester_id=semester_id,
        department_id=department_id,
        programme_id=programme_id,
        search=search,
    )

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Outstanding Fees"

    headers = [
        "Student",
        "Admission No.",
        "Programme",
        "Department",
        "Invoices",
        "Gross Invoiced",
        "Cash Collections",
        "Credits Applied",
        "Outstanding",
    ]

    worksheet.append(
        headers
    )

    for cell in worksheet[1]:

        cell.font = Font(
            bold=True
        )

    # --------------------------------------------------------
    # STUDENT ROWS
    # --------------------------------------------------------

    for row in report["rows"]:

        worksheet.append(
            [
                row.get(
                    "student_name",
                    "",
                ),

                row.get(
                    "admission_number",
                    "",
                ),

                row.get(
                    "programme",
                    "",
                ),

                row.get(
                    "department",
                    "",
                ),

                row.get(
                    "invoice_count",
                    0,
                ),

                float(
                    row.get(
                        "gross_invoiced",
                        0,
                    )
                    or 0
                ),

                float(
                    row.get(
                        "cash_collections",
                        0,
                    )
                    or 0
                ),

                float(
                    row.get(
                        "credits_applied",
                        0,
                    )
                    or 0
                ),

                float(
                    row.get(
                        "outstanding",
                        0,
                    )
                    or 0
                ),
            ]
        )

    # --------------------------------------------------------
    # TOTAL
    # --------------------------------------------------------

    worksheet.append([])

    summary = report[
        "summary"
    ]

    worksheet.append(
        [
            "TOTAL",
            "",
            "",
            "",

            summary.get(
                "invoice_count",
                0,
            ),

            float(
                summary.get(
                    "gross_invoiced",
                    0,
                )
                or 0
            ),

            float(
                summary.get(
                    "cash_collections",
                    0,
                )
                or 0
            ),

            float(
                summary.get(
                    "credits_applied",
                    0,
                )
                or 0
            ),

            float(
                summary.get(
                    "outstanding",
                    0,
                )
                or 0
            ),
        ]
    )

    # --------------------------------------------------------
    # COLUMN WIDTHS
    # --------------------------------------------------------

    for column in worksheet.columns:

        max_length = 0

        column_letter = (
            column[0].column_letter
        )

        for cell in column:

            value = str(
                cell.value
                or ""
            )

            if len(value) > max_length:

                max_length = len(
                    value
                )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            40,
        )

    # --------------------------------------------------------
    # RESPONSE
    # --------------------------------------------------------

    output = io.BytesIO()

    workbook.save(
        output
    )

    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )

    response[
        "Content-Disposition"
    ] = (
        'attachment; '
        'filename="XORADEX_Outstanding_Fees.xlsx"'
    )

    return response


# ============================================================
# OUTSTANDING FEES PDF
# ============================================================

@login_required
def outstanding_fees_pdf(request):

    semester_id = (
        request.GET.get(
            "semester"
        )
        or None
    )

    department_id = (
        request.GET.get(
            "department"
        )
        or None
    )

    programme_id = (
        request.GET.get(
            "programme"
        )
        or None
    )

    search = (
        request.GET.get(
            "search"
        )
        or ""
    ).strip()

    report = get_outstanding_fees_report(
        request.user,
        semester_id=semester_id,
        department_id=department_id,
        programme_id=programme_id,
        search=search,
    )

    output = io.BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=25,
        leftMargin=25,
        topMargin=25,
        bottomMargin=25,
    )

    styles = getSampleStyleSheet()

    title_style = styles[
        "Title"
    ]

    heading_style = styles[
        "Heading2"
    ]

    normal_style = styles[
        "Normal"
    ]

    elements = []

    # --------------------------------------------------------
    # TITLE
    # --------------------------------------------------------

    elements.append(
        Paragraph(
            "XORADEX EDUCORE",
            title_style,
        )
    )

    elements.append(
        Paragraph(
            "Outstanding Fees Report",
            heading_style,
        )
    )

    # --------------------------------------------------------
    # SEMESTER
    # --------------------------------------------------------

    semester = report.get(
        "semester"
    )

    if semester:

        semester_name = str(
            semester
        )

        elements.append(
            Paragraph(
                f"<b>Semester:</b> "
                f"{semester_name}",
                normal_style,
            )
        )

    # --------------------------------------------------------
    # FILTER CONTEXT
    # --------------------------------------------------------

    if department_id:

        elements.append(
            Paragraph(
                "<b>Department:</b> "
                "Filtered",
                normal_style,
            )
        )

    if programme_id:

        elements.append(
            Paragraph(
                "<b>Programme:</b> "
                "Filtered",
                normal_style,
            )
        )

    if search:

        elements.append(
            Paragraph(
                f"<b>Search:</b> "
                f"{search}",
                normal_style,
            )
        )

    # --------------------------------------------------------
    # SUMMARY
    # --------------------------------------------------------

    summary = report[
        "summary"
    ]

    elements.append(
        Paragraph(
            f"<b>Students with Outstanding Fees:</b> "
            f"{summary.get('student_count', 0)}",
            normal_style,
        )
    )

    elements.append(
        Paragraph(
            f"<b>Outstanding Invoices:</b> "
            f"{summary.get('invoice_count', 0)}",
            normal_style,
        )
    )

    elements.append(
        Paragraph(
            f"<b>Gross Invoiced:</b> "
            f"KSh {float(summary.get('gross_invoiced', 0) or 0):,.2f}",
            normal_style,
        )
    )

    elements.append(
        Paragraph(
            f"<b>Cash Collections:</b> "
            f"KSh {float(summary.get('cash_collections', 0) or 0):,.2f}",
            normal_style,
        )
    )

    elements.append(
        Paragraph(
            f"<b>Credits Applied:</b> "
            f"KSh {float(summary.get('credits_applied', 0) or 0):,.2f}",
            normal_style,
        )
    )

    elements.append(
        Paragraph(
            f"<b>Total Outstanding:</b> "
            f"KSh {float(summary.get('outstanding', 0) or 0):,.2f}",
            normal_style,
        )
    )

    elements.append(
        Spacer(
            1,
            8,
        )
    )

    # --------------------------------------------------------
    # TABLE
    # --------------------------------------------------------

    data = [
        [
            "Student",
            "Admission No.",
            "Programme",
            "Department",
            "Invoices",
            "Gross Invoiced",
            "Cash",
            "Credits",
            "Outstanding",
        ]
    ]

    for row in report["rows"]:

        data.append(
            [
                str(
                    row.get(
                        "student_name",
                        "",
                    )
                ),

                str(
                    row.get(
                        "admission_number",
                        "",
                    )
                ),

                str(
                    row.get(
                        "programme",
                        "",
                    )
                ),

                str(
                    row.get(
                        "department",
                        "",
                    )
                ),

                str(
                    row.get(
                        "invoice_count",
                        0,
                    )
                ),

                f"KSh {float(row.get('gross_invoiced', 0) or 0):,.2f}",

                f"KSh {float(row.get('cash_collections', 0) or 0):,.2f}",

                f"KSh {float(row.get('credits_applied', 0) or 0):,.2f}",

                f"KSh {float(row.get('outstanding', 0) or 0):,.2f}",
            ]
        )

    # --------------------------------------------------------
    # TOTAL ROW
    # --------------------------------------------------------

    data.append(
        [
            "TOTAL",
            "",
            "",
            "",

            str(
                summary.get(
                    "invoice_count",
                    0,
                )
            ),

            f"KSh {float(summary.get('gross_invoiced', 0) or 0):,.2f}",

            f"KSh {float(summary.get('cash_collections', 0) or 0):,.2f}",

            f"KSh {float(summary.get('credits_applied', 0) or 0):,.2f}",

            f"KSh {float(summary.get('outstanding', 0) or 0):,.2f}",
        ]
    )

    table = Table(
        data,
        repeatRows=1,
    )

    table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.grey,
                ),

                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white,
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),

                (
                    "FONTNAME",
                    (0, -1),
                    (-1, -1),
                    "Helvetica-Bold",
                ),

                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey,
                ),

                (
                    "ALIGN",
                    (4, 1),
                    (-1, -1),
                    "RIGHT",
                ),

                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),

                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),

                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),

                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
            ]
        )
    )

    elements.append(
        table
    )

    # --------------------------------------------------------
    # BUILD
    # --------------------------------------------------------

    document.build(
        elements
    )

    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type="application/pdf",
    )

    response[
        "Content-Disposition"
    ] = (
        'attachment; '
        'filename="XORADEX_Outstanding_Fees.pdf"'
    )

    return response


# ============================================================
# EXAMINATION PERFORMANCE
# ============================================================

@login_required
def examination_performance(request):

    semester_id = (
        request.GET.get(
            "semester"
        )
        or None
    )

    department_id = (
        request.GET.get(
            "department"
        )
        or None
    )

    programme_id = (
        request.GET.get(
            "programme"
        )
        or None
    )

    search = (
        request.GET.get(
            "search"
        )
        or ""
    ).strip()

    report = get_examination_performance_report(
        semester_id=semester_id,
        department_id=department_id,
        programme_id=programme_id,
        search=search,
    )

    semesters = (
        get_all_semesters(
            request.user
        )
        .order_by(
            "-academic_year__year_name",
            "-start_date",
            "-id",
        )
    )

    departments = (
        Department.objects
        .filter(
            is_active=True
        )
        .order_by(
            "name"
        )
    )

    programmes = (
        Programme.objects
        .filter(
            is_active=True
        )
        .select_related(
            "course",
            "course__department",
        )
        .order_by(
            "name"
        )
    )

    # --------------------------------------------------------
    # HOD RESTRICTION
    # --------------------------------------------------------

    if request.user.groups.filter(
        name="HOD"
    ).exists():

        departments = departments.filter(
            hod=request.user
        )

        programmes = programmes.filter(
            course__department__hod=request.user
        )

    # --------------------------------------------------------
    # RENDER
    # --------------------------------------------------------

    return render(
        request,
        "reports/examination_performance.html",
        {
            "report":
                report,

            "rows":
                report["rows"],

            "summary":
                report["summary"],

            "semesters":
                semesters,

            "departments":
                departments,

            "programmes":
                programmes,

            "selected_semester":
                semester_id,

            "selected_department":
                department_id,

            "selected_programme":
                programme_id,

            "search":
                search,

            "report_role":
                get_report_role(
                    request.user
                ),
        },
    )

# ============================================================
# MISSING MARKS
# ============================================================

@login_required
def missing_marks(request):

    semester_id = (
        request.GET.get(
            "semester"
        )
        or None
    )

    department_id = (
        request.GET.get(
            "department"
        )
        or None
    )

    programme_id = (
        request.GET.get(
            "programme"
        )
        or None
    )

    search = (
        request.GET.get(
            "search"
        )
        or ""
    ).strip()

    report = get_missing_marks_report(
        semester_id=semester_id,
        department_id=department_id,
        programme_id=programme_id,
        search=search,
    )

    semesters = (
        get_all_semesters(
            request.user
        )
        .order_by(
            "-academic_year__year_name",
            "-start_date",
            "-id",
        )
    )

    departments = (
        Department.objects
        .filter(
            is_active=True
        )
        .order_by(
            "name"
        )
    )

    programmes = (
        Programme.objects
        .filter(
            is_active=True
        )
        .select_related(
            "course",
            "course__department",
        )
        .order_by(
            "name"
        )
    )

    # --------------------------------------------------------
    # HOD RESTRICTION
    # --------------------------------------------------------

    if request.user.groups.filter(
        name="HOD"
    ).exists():

        departments = departments.filter(
            hod=request.user
        )

        programmes = programmes.filter(
            course__department__hod=request.user
        )

    return render(
        request,
        "reports/missing_marks.html",
        {
            "rows":
                report["rows"],

            "count":
                report["count"],

            "semesters":
                semesters,

            "departments":
                departments,

            "programmes":
                programmes,

            "selected_semester":
                semester_id,

            "selected_department":
                department_id,

            "selected_programme":
                programme_id,

            "search":
                search,

            "report_role":
                get_report_role(
                    request.user
                ),
        },
    )


# ============================================================
# EXAMINATION PERFORMANCE - EXCEL
# ============================================================

@login_required
def examination_performance_excel(request):

    semester_id = (
        request.GET.get("semester")
        or None
    )

    department_id = (
        request.GET.get("department")
        or None
    )

    programme_id = (
        request.GET.get("programme")
        or None
    )

    search = (
        request.GET.get("search")
        or ""
    ).strip()

    report = get_examination_performance_report(
        semester_id=semester_id,
        department_id=department_id,
        programme_id=programme_id,
        search=search,
    )

    # --------------------------------------------------------
    # IMPORTS
    # --------------------------------------------------------

    from openpyxl import Workbook
    from django.http import HttpResponse

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Examination Performance"

    # --------------------------------------------------------
    # HEADINGS
    # --------------------------------------------------------

    worksheet.append(
        [
            "Admission No.",
            "Student",
            "Department",
            "Programme",
            "Unit Code",
            "Unit",
            "CAT 1",
            "CAT 2",
            "Exam",
            "Total",
            "Grade",
            "Remarks",
        ]
    )

    # --------------------------------------------------------
    # DATA
    # --------------------------------------------------------

    for row in report["rows"]:

        worksheet.append(
            [
                row["admission_no"],
                row["student_name"],
                row["department"],
                row["programme"],
                row["unit_code"],
                row["unit_name"],
                row["cat1"],
                row["cat2"],
                row["exam"],
                row["total"],
                row["grade"],
                row["remarks"],
            ]
        )

    # --------------------------------------------------------
    # RESPONSE
    # --------------------------------------------------------

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response[
        "Content-Disposition"
    ] = (
        'attachment; filename="'
        'examination_performance.xlsx"'
    )

    workbook.save(
        response
    )

    return response


# ============================================================
# EXAMINATION PERFORMANCE - PDF
# ============================================================

@login_required
def examination_performance_pdf(request):

    semester_id = (
        request.GET.get("semester")
        or None
    )

    department_id = (
        request.GET.get("department")
        or None
    )

    programme_id = (
        request.GET.get("programme")
        or None
    )

    search = (
        request.GET.get("search")
        or ""
    ).strip()

    report = get_examination_performance_report(
        semester_id=semester_id,
        department_id=department_id,
        programme_id=programme_id,
        search=search,
    )

    response = HttpResponse(
        content_type="application/pdf"
    )

    response[
        "Content-Disposition"
    ] = (
        'attachment; filename="'
        'examination_performance.pdf"'
    )

    document = SimpleDocTemplate(
        response,
        pagesize=landscape(A4),
        rightMargin=20,
        leftMargin=20,
        topMargin=20,
        bottomMargin=20,
    )

    styles = getSampleStyleSheet()

    elements = []

    elements.append(
        Paragraph(
            "Examination Performance",
            styles["Title"],
        )
    )

    elements.append(
        Spacer(
            1,
            10,
        )
    )

    elements.append(
        Paragraph(
            (
                f"Students: "
                f"{report['summary']['student_count']} "
                f"&nbsp;&nbsp; "
                f"Results: "
                f"{report['summary']['result_count']} "
                f"&nbsp;&nbsp; "
                f"Passed: "
                f"{report['summary']['passed']} "
                f"&nbsp;&nbsp; "
                f"Failed: "
                f"{report['summary']['failed']} "
                f"&nbsp;&nbsp; "
                f"Pass Rate: "
                f"{report['summary']['pass_rate']}%"
            ),
            styles["Normal"],
        )
    )

    elements.append(
        Spacer(
            1,
            10,
        )
    )

    data = [
        [
            "Admission No.",
            "Student",
            "Department",
            "Programme",
            "Unit",
            "CAT 1",
            "CAT 2",
            "Exam",
            "Total",
            "Grade",
            "Remarks",
        ]
    ]

    for row in report["rows"]:

        data.append(
            [
                row["admission_no"],
                row["student_name"],
                row["department"],
                row["programme"],
                row["unit_code"],
                str(row["cat1"] or "-"),
                str(row["cat2"] or "-"),
                str(row["exam"] or "-"),
                str(row["total"]),
                row["grade"],
                row["remarks"],
            ]
        )

    table = Table(
        data,
        repeatRows=1,
    )

    table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey,
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
            ]
        )
    )

    elements.append(
        table
    )

    document.build(
        elements
    )

    return response



# ============================================================
# STUDENT RESULT HISTORY
# ============================================================
@login_required
def student_result_history(request, student_id):

    from collections import OrderedDict

    from django.shortcuts import get_object_or_404

    from students.models import Student, Result

    student = get_object_or_404(
        Student,
        id=student_id,
    )

    results = list(
        Result.objects
        .filter(
            enrollment__student_id=student_id,
        )
        .select_related(
            "enrollment",
            "enrollment__student",
            "enrollment__academic_year",
            "enrollment__semester",
            "enrollment__programme_level",
            "registration",
            "registration__unit_offering",
            "registration__unit_offering__unit",
            "unit_offering",
            "unit_offering__unit",
        )
        .order_by(
            "enrollment__academic_year__year_name",
            "enrollment__semester__start_date",
            "enrollment__programme_level__progression_order",
            "unit_offering__unit__code",
        )
    )

    # ============================================================
    # GROUP RESULTS BY ENROLLMENT / SEMESTER
    # ============================================================

    grouped = OrderedDict()

    for result in results:

        enrollment = result.enrollment

        key = enrollment.id

        if key not in grouped:
            grouped[key] = {
                "enrollment": enrollment,
                "results": [],
            }

        grouped[key]["results"].append(result)

    grouped_results = list(grouped.values())

    # ============================================================
    # SUMMARY
    # ============================================================

    total_results = len(results)

    passed = sum(
        1
        for result in results
        if result.remarks == "PASS"
    )

    failed = sum(
        1
        for result in results
        if result.remarks == "FAIL"
    )

    return render(
        request,
        "reports/student_result_history.html",
        {
            "student": student,
            "results": results,
            "grouped_results": grouped_results,
            "total_results": total_results,
            "passed": passed,
            "failed": failed,
        },
    )


# ============================================================
# STAFF REGISTER REPORT
# ============================================================

@login_required
def staff_register(request):
    """Institutional staff register based on existing user accounts."""

    role = get_report_role(request.user)

    if role not in {
        "ADMIN",
        "PRINCIPAL",
        "REGISTRAR",
    }:
        return HttpResponse(
            "You are not authorised to access the Staff Register.",
            status=403,
        )

    selected_role = request.GET.get("role", "")
    selected_status = request.GET.get("status", "")
    search = request.GET.get("search", "")

    report = get_staff_register_report(
        search=search,
        role=selected_role or None,
        status=selected_status or None,
    )

    return render(
        request,
        "reports/staff_register.html",
        {
            **report,
            "staff_roles": get_staff_roles(),
            "selected_role": selected_role,
            "selected_status": selected_status,
            "search": search,
            "report_role": role,
        },
    )


@login_required
def staff_register_excel(request):
    """Download the institutional Staff Register as Excel."""

    role = get_report_role(request.user)

    if role not in {
        "ADMIN",
        "PRINCIPAL",
        "REGISTRAR",
    }:
        return HttpResponse(
            "You are not authorised to access the Staff Register.",
            status=403,
        )

    report = get_staff_register_report(
        search=request.GET.get("search", ""),
        role=request.GET.get("role") or None,
        status=request.GET.get("status") or None,
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Staff Register"

    headers = [
        "Name",
        "Username",
        "Email",
        "Role(s)",
        "Status",
        "Django Staff",
        "Date Joined",
        "Last Login",
    ]

    sheet.append(headers)

    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for row in report["rows"]:
        sheet.append(
            [
                row["name"],
                row["username"],
                row["email"],
                row["roles"],
                row["status"],
                "Yes" if row["is_staff"] else "No",
                row["date_joined"].strftime("%Y-%m-%d") if row["date_joined"] else "",
                row["last_login"].strftime("%Y-%m-%d %H:%M") if row["last_login"] else "Never",
            ]
        )

    for column in sheet.columns:
        maximum = max(
            len(str(cell.value or ""))
            for cell in column
        )
        sheet.column_dimensions[column[0].column_letter].width = min(
            maximum + 2,
            40,
        )

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )

    response["Content-Disposition"] = (
        'attachment; filename="staff_register.xlsx"'
    )

    return response


# ============================================================
# USER ACCOUNTS REPORT
# ============================================================

@login_required
def user_accounts(request):

    if not request.user.is_superuser:
        return render(
            request,
            "403.html",
            status=403,
        )

    search = request.GET.get(
        "search",
        "",
    ).strip()

    group_name = request.GET.get(
        "group",
        "",
    ).strip()

    status = request.GET.get(
        "status",
        "",
    ).strip()

    report = get_user_accounts_report(
        search=search,
        group_name=group_name,
        status=status,
    )

    User = get_user_model()

    groups = (
        User.groups.field.remote_field.model.objects
        .order_by("name")
    )

    return render(
        request,
        "reports/user_accounts.html",
        {
            "rows": report["rows"],
            "count": report["count"],
            "groups": groups,
            "search": search,
            "selected_group": group_name,
            "selected_status": status,
        },
    )


# ============================================================
# AUDIT TRAIL REPORT
# ============================================================

@login_required
def audit_trail(request):

    if not request.user.is_superuser:
        return render(
            request,
            "403.html",
            status=403,
        )

    search = request.GET.get(
        "search",
        "",
    ).strip()

    module = request.GET.get(
        "module",
        "",
    ).strip()

    action = request.GET.get(
        "action",
        "",
    ).strip()

    severity = request.GET.get(
        "severity",
        "",
    ).strip()

    user_id = request.GET.get(
        "user",
        "",
    ).strip()

    report = get_audit_trail_report(
        search=search,
        module=module,
        action=action,
        severity=severity,
        user_id=user_id,
    )

    User = get_user_model()

    users = (
        User.objects
        .filter(
            activity_logs__isnull=False
        )
        .distinct()
        .order_by(
            "first_name",
            "last_name",
            "username",
        )
    )

    modules = [
        choice[0]
        for choice in ActivityLog.MODULE_CHOICES
    ]

    actions = [
        choice[0]
        for choice in ActivityLog.ACTION_CHOICES
    ]

    severities = [
        choice[0]
        for choice in ActivityLog.SEVERITY_CHOICES
    ]

    return render(
        request,
        "reports/audit_trail.html",
        {
            "rows": report["rows"],
            "count": report["count"],
            "users": users,
            "modules": modules,
            "actions": actions,
            "severities": severities,
            "search": search,
            "selected_module": module,
            "selected_action": action,
            "selected_severity": severity,
            "selected_user": user_id,
        },
    )